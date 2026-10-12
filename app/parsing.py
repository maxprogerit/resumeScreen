"""Text extraction from uploaded resumes.

File type is decided by content (magic bytes / archive structure), never by the filename
or the client-supplied content type. Extracted text is plain text; nothing from an
uploaded file is ever rendered as HTML.
"""

import io
import re
import threading
import unicodedata
import zipfile
from dataclasses import dataclass, field
from enum import StrEnum
from functools import lru_cache
from typing import Any

import pdfplumber
from docx import Document
from pypdf import PdfReader
from pypdf.errors import PyPdfError

MAX_TEXT_CHARS = 100_000
MAX_PDF_PAGES = 30
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 1000
MAX_OCR_PAGES = 5
MAX_OCR_PIXELS = 2600  # longest side of the rendered page image
MIN_TEXT_CHARS = 20

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class FileKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"


class UnsupportedFileType(Exception):
    pass


class ParseError(Exception):
    """The file is a supported type but no usable text could be extracted.

    The message is a short machine-readable reason, never document content.
    """


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    # How the text was obtained: "ocr" (scanned PDF, no text layer) and
    # "hidden_text_removed" (white / microscopic / vanish text was dropped).
    notes: list[str] = field(default_factory=list)


def sniff_file_kind(data: bytes) -> FileKind:
    if data.startswith(b"%PDF-"):
        return FileKind.PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" in zf.namelist() and "[Content_Types].xml" in zf.namelist():
                    return FileKind.DOCX
        except zipfile.BadZipFile:
            pass
    raise UnsupportedFileType("only PDF and DOCX files are supported")


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = _CONTROL.sub("", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


# ---------------------------------------------------------------- PDF


def _is_white(color: Any) -> bool:
    if color is None:
        return False
    values = color if isinstance(color, tuple | list) else (color,)
    return len(values) > 0 and all(isinstance(v, int | float) and v >= 0.99 for v in values)


def _hidden_char(char: dict[str, Any]) -> bool:
    """White or microscopic text: invisible to a human reader, visible to a parser."""
    return _is_white(char.get("non_stroking_color")) or float(char.get("size") or 10) < 1.5


def _pdf_text(data: bytes) -> tuple[str, bool]:
    """Returns (text, hidden_text_removed)."""
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ParseError("pdf_encrypted")
        pages = reader.pages[:MAX_PDF_PAGES]
        plain = "\n".join(page.extract_text() or "" for page in pages)
    except PyPdfError as exc:
        raise ParseError("pdf_unreadable") from exc
    except (ValueError, KeyError, TypeError, RecursionError) as exc:
        raise ParseError("pdf_unreadable") from exc

    try:
        with pdfplumber.open(io.BytesIO(data)) as doc:
            pdf_pages = doc.pages[:MAX_PDF_PAGES]
            if not any(_hidden_char(c) for page in pdf_pages for c in page.chars):
                return plain, False
            visible = [
                page.filter(
                    lambda o: o.get("object_type") != "char" or not _hidden_char(o)
                ).extract_text()
                for page in pdf_pages
            ]
        return "\n".join(visible), True
    except Exception:  # noqa: BLE001 - screening is best effort and must never block parsing
        return plain, False


# ---------------------------------------------------------------- OCR (scanned PDFs)

_ocr_lock = threading.Lock()


@lru_cache(maxsize=1)
def _ocr_engine() -> Any:
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR()


def _ocr_pdf(data: bytes) -> str:
    """OCR the first pages of a PDF that has no text layer."""
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(data)
        engine = _ocr_engine()
        lines: list[str] = []
        for index in range(min(len(pdf), MAX_OCR_PAGES)):
            page = pdf[index]
            width, height = page.get_size()
            scale = min(2.0, MAX_OCR_PIXELS / max(width, height, 1.0))
            image = page.render(scale=scale).to_numpy()
            with _ocr_lock:  # the ONNX session is shared between threads
                result, _ = engine(image)
            lines += [item[1] for item in (result or [])]
        return "\n".join(lines)
    except Exception as exc:  # noqa: BLE001
        raise ParseError("ocr_failed") from exc


# ---------------------------------------------------------------- DOCX


def _run_hidden(run: Any) -> bool:
    if run.font.hidden:
        return True
    try:
        color = run.font.color
        return color is not None and color.rgb is not None and str(color.rgb).upper() == "FFFFFF"
    except (ValueError, AttributeError):
        return False


def _paragraph_text(paragraph: Any) -> tuple[str, bool]:
    runs = paragraph.runs
    if any(_run_hidden(r) for r in runs):
        return "".join(r.text for r in runs if not _run_hidden(r)), True
    return str(paragraph.text), False


def _docx_text(data: bytes) -> tuple[str, bool]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if (
                len(infos) > MAX_DOCX_ENTRIES
                or sum(i.file_size for i in infos) > MAX_DOCX_UNCOMPRESSED
            ):
                raise ParseError("docx_too_large")
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ParseError("docx_unreadable") from exc
    parts: list[str] = []
    removed = False
    for paragraph in doc.paragraphs:
        text, hid = _paragraph_text(paragraph)
        parts.append(text)
        removed = removed or hid
    for table in doc.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                cell_parts: list[str] = []
                for paragraph in cell.paragraphs:
                    text, hid = _paragraph_text(paragraph)
                    cell_parts.append(text)
                    removed = removed or hid
                cells.append(" ".join(cell_parts).strip())
            parts.append(" | ".join(cells))
    return "\n".join(parts), removed


# ---------------------------------------------------------------- public API


def extract_document(data: bytes) -> ParsedDocument:
    """Normalised plain text plus notes on how it was obtained.

    Raises UnsupportedFileType or ParseError. Scanned PDFs (no text layer) go through OCR.
    """
    kind = sniff_file_kind(data)
    notes: list[str] = []
    if kind is FileKind.PDF:
        raw, hidden = _pdf_text(data)
        text = normalize_text(raw)
        if len(text) < MIN_TEXT_CHARS:
            text = normalize_text(_ocr_pdf(data))
            notes.append("ocr")
    else:
        raw, hidden = _docx_text(data)
        text = normalize_text(raw)
    if hidden:
        notes.append("hidden_text_removed")
    if len(text) < MIN_TEXT_CHARS:
        raise ParseError("no_extractable_text")
    return ParsedDocument(text=text, notes=notes)


def extract_text(data: bytes) -> str:
    return extract_document(data).text


def safe_filename(name: str | None) -> str:
    """Display-safe filename: basename only, control characters removed, length capped."""
    base = (name or "resume").replace("\\", "/").rsplit("/", 1)[-1]
    base = _CONTROL.sub("", base).strip()
    return (base or "resume")[:255]
